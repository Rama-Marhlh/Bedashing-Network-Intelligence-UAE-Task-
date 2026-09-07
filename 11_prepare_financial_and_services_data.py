"""Prepare service-price and theoretical branch-capacity JSON datasets.

Capacity values are scenarios, not actual revenue, and are never used to alter
branch recommendations.
"""

from __future__ import annotations

import json
import logging
import math
import os
import statistics
import tempfile
from collections.abc import Iterable, Mapping
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = (
    ROOT / "data" / "external" / "Bedashing_Official_Services_and_Price_Intelligence.xlsx"
)
DEFAULT_OUTPUT_DIR = ROOT / "app_data"
REQUIRED_SHEETS = {
    "Official Services",
    "Service Summary",
    "Branches 24",
    "Capacity Model",
    "Sources & Audit",
    "Checks",
}
SERVICE_FIELDS = (
    "service_id",
    "category",
    "subcategory",
    "service_name_official",
    "variant",
    "duration_minutes",
    "price_aed",
    "price_basis",
    "tax_status",
    "source_url",
    "retrieved_at",
    "source_type",
    "quality_flag",
    "notes",
)
OPERATING_DAYS = 30
LIST_PRODUCTIVITY = 329.0
REALIZATION_FACTOR = 0.90
DISCLAIMER = "Scenario-based theoretical capacity, not actual branch revenue."

LOGGER = logging.getLogger(__name__)


@dataclass(frozen=True)
class Scenario:
    productive_staff: int
    utilization_rate: float


SCENARIOS = {
    "conservative": Scenario(6, 0.50),
    "base": Scenario(10, 0.65),
    "high": Scenario(14, 0.80),
}


@dataclass(frozen=True)
class Branch:
    official_store_id: int | str
    branch_name: str
    average_open_hours_per_day: float
    opening_hours_source: str


def hourly_productivity(
    price_aed: float,
    duration_minutes: float | None,
    price_basis: str,
    quality_flag: str,
) -> float | None:
    """Return list-price productivity only for complete, approved service rows."""
    if duration_minutes is None or price_basis != "per service" or quality_flag != "OK":
        return None
    if duration_minutes <= 0:
        raise ValueError("duration_minutes must be positive when provided")
    return price_aed * 60 / duration_minutes


def monthly_capacity(
    opening_hours_per_day: float,
    operating_days_per_month: int,
    productive_staff: int,
    utilization_rate: float,
    list_productivity_aed_per_hour: float,
    realization_factor: float,
) -> float:
    """Calculate theoretical gross service-sales capacity."""
    values = (
        opening_hours_per_day,
        operating_days_per_month,
        productive_staff,
        utilization_rate,
        list_productivity_aed_per_hour,
        realization_factor,
    )
    if any(value < 0 for value in values):
        raise ValueError("capacity inputs cannot be negative")
    return math.prod(values)


def validate_branches(branches: Iterable[Branch], expected_count: int = 24) -> list[Branch]:
    result = list(branches)
    names = [branch.branch_name.casefold() for branch in result]
    ids = [str(branch.official_store_id).casefold() for branch in result]
    if len(names) != len(set(names)) or len(ids) != len(set(ids)):
        raise ValueError("duplicate branch detected (branch name or official_store_id)")
    if len(result) != expected_count:
        raise ValueError(f"branch count must be {expected_count}; found {len(result)}")
    return result


def _rows_as_dicts(sheet: Any, header_row: int = 1) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < header_row:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[header_row - 1]]
    return [
        dict(zip(headers, row, strict=False))
        for row in rows[header_row:]
        if any(value is not None for value in row)
    ]


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def read_services(sheet: Any) -> list[dict[str, Any]]:
    rows = _rows_as_dicts(sheet)
    services = [{field: _json_value(row.get(field)) for field in SERVICE_FIELDS} for row in rows]
    if not services:
        raise ValueError("service variants are missing")
    for index, service in enumerate(services, start=2):
        price = service["price_aed"]
        if (
            isinstance(price, bool)
            or not isinstance(price, (int, float))
            or not math.isfinite(price)
            or price <= 0
        ):
            raise ValueError(f"invalid price_aed in Official Services row {index}: {price!r}")
    return services


def summarize_services(services: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for service in services:
        grouped.setdefault(str(service["category"]), []).append(service)
    summaries = []
    for category in sorted(grouped):
        rows = grouped[category]
        prices = [float(row["price_aed"]) for row in rows]
        productivities = [
            value
            for row in rows
            if (
                value := hourly_productivity(
                    float(row["price_aed"]),
                    float(row["duration_minutes"]) if row["duration_minutes"] is not None else None,
                    str(row["price_basis"]),
                    str(row["quality_flag"]),
                )
            )
            is not None
        ]
        summaries.append(
            {
                "category": category,
                "variant_count": len(rows),
                "minimum_price_aed": min(prices),
                "maximum_price_aed": max(prices),
                "median_price_aed": statistics.median(prices),
                "average_price_aed": round(statistics.fmean(prices), 2),
                "median_productivity_aed_per_hour": (
                    round(statistics.median(productivities), 2) if productivities else None
                ),
            }
        )
    return summaries


def read_branches(branch_sheet: Any, capacity_sheet: Any) -> list[Branch]:
    capacity_rows = _rows_as_dicts(capacity_sheet, header_row=16)
    capacity_sources = {
        str(row["Branch"]): str(row["Source status"])
        for row in capacity_rows
        if row.get("Branch") is not None
    }
    branches = []
    for row in _rows_as_dicts(branch_sheet):
        name = str(row.get("official_branch_name") or "").strip()
        hours = row.get("derived_avg_open_hours_day")
        if not name or not isinstance(hours, (int, float)) or hours < 0:
            raise ValueError(f"invalid branch name/opening hours: {name!r}, {hours!r}")
        branches.append(
            Branch(
                official_store_id=row["official_store_id"],
                branch_name=name,
                average_open_hours_per_day=float(hours),
                opening_hours_source=capacity_sources.get(name, "Google Places opening hours"),
            )
        )
    return validate_branches(branches)


def build_capacity_records(branches: list[Branch]) -> list[dict[str, Any]]:
    records = []
    for branch in branches:
        values = {
            name: round(
                monthly_capacity(
                    branch.average_open_hours_per_day,
                    OPERATING_DAYS,
                    scenario.productive_staff,
                    scenario.utilization_rate,
                    LIST_PRODUCTIVITY,
                    REALIZATION_FACTOR,
                ),
                2,
            )
            for name, scenario in SCENARIOS.items()
        }
        if any(value < 0 for value in values.values()):
            raise ValueError(f"negative scenario output for {branch.branch_name}")
        records.append(
            {
                **asdict(branch),
                "conservative_monthly_capacity_aed": values["conservative"],
                "base_monthly_capacity_aed": values["base"],
                "high_monthly_capacity_aed": values["high"],
                "financial_confidence": "LOW",
                "estimate_type": "THEORETICAL_GROSS_SERVICE_SALES_CAPACITY",
                "use_in_branch_decision": False,
            }
        )
    return records


def build_revenue_explanations(branches: list[Branch]) -> list[dict[str, Any]]:
    scenario = SCENARIOS["base"]
    records = []
    for branch in branches:
        available = branch.average_open_hours_per_day * OPERATING_DAYS * scenario.productive_staff
        productive = available * scenario.utilization_rate
        realized = LIST_PRODUCTIVITY * REALIZATION_FACTOR
        estimate = monthly_capacity(
            branch.average_open_hours_per_day,
            OPERATING_DAYS,
            scenario.productive_staff,
            scenario.utilization_rate,
            LIST_PRODUCTIVITY,
            REALIZATION_FACTOR,
        )
        records.append(
            {
                "branch_name": branch.branch_name,
                "base_estimate_aed": round(estimate, 2),
                "inputs": {
                    "opening_hours_per_day": branch.average_open_hours_per_day,
                    "operating_days_per_month": OPERATING_DAYS,
                    "productive_staff": scenario.productive_staff,
                    "utilization_rate": scenario.utilization_rate,
                    "list_productivity_aed_per_hour": int(LIST_PRODUCTIVITY),
                    "realization_factor": REALIZATION_FACTOR,
                },
                "calculation": {
                    "available_staff_hours": round(available, 2),
                    "productive_staff_hours": round(productive, 2),
                    "realized_productivity_aed_per_hour": round(realized, 2),
                    "estimated_monthly_capacity_aed": round(estimate, 2),
                },
                "explanation": (
                    "Base theoretical capacity uses "
                    f"{branch.average_open_hours_per_day:g} open hours/day, "
                    f"{scenario.productive_staff} productive staff, "
                    f"{scenario.utilization_rate:.0%} "
                    "utilization, and realized list-price productivity."
                ),
                "disclaimer": DISCLAIMER,
            }
        )
    return records


def build_methodology() -> dict[str, Any]:
    return {
        "estimate_type": "THEORETICAL_GROSS_SERVICE_SALES_CAPACITY",
        "formula": (
            "open hours/day × operating days/month × productive staff × utilization rate × "
            "list productivity AED/hour × realization factor"
        ),
        "sourced_inputs": {
            "opening_hours_per_day": "Derived from Google Places opening hours in Branches 24.",
            "service_prices_and_durations": (
                "Official Bedashing service pages; source URLs and quality flags retained "
                "per service variant."
            ),
        },
        "assumed_inputs": {
            "operating_days_per_month": OPERATING_DAYS,
            "list_productivity_aed_per_hour": int(LIST_PRODUCTIVITY),
            "realization_factor": REALIZATION_FACTOR,
            "productive_staff_and_utilization": (
                "Scenario assumptions from Capacity Model; not observed branch operations."
            ),
        },
        "scenario_definitions": {
            name: {
                "productive_staff": value.productive_staff,
                "utilization_rate": value.utilization_rate,
            }
            for name, value in SCENARIOS.items()
        },
        "limitations": [
            "Capacity is not actual revenue, profit, cash flow, or a forecast.",
            "List-price productivity does not represent transaction mix, discounts, VAT "
            "normalization, or actual realized tickets.",
            "Service availability, staffing, utilization, and operating constraints may vary "
            "by branch.",
        ],
        "missing_internal_data": [
            "actual branch sales and transactions",
            "bookings, cancellations, and utilization",
            "staff headcount, schedules, skills, and productivity",
            "chairs, rooms, and service capacity constraints",
            "discounts, packages, service mix, and realized ticket values",
        ],
        "why_capacity_estimates_are_excluded_from_PROTECT_HOLD_SHRINK": (
            "They combine public factual inputs with unvalidated operating assumptions. "
            "Using them in recommendation scores could misstate branch performance and modify "
            "decisions without actual internal data."
        ),
        "use_in_branch_decision": False,
    }


def atomic_write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(
                payload, handle, ensure_ascii=False, indent=2, sort_keys=False, allow_nan=False
            )
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_name, path)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise


def prepare(
    workbook_path: Path = DEFAULT_WORKBOOK, output_dir: Path = DEFAULT_OUTPUT_DIR
) -> Mapping[str, int]:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"required workbook not found: {workbook_path}")
    LOGGER.info("Reading workbook: %s", workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        missing = sorted(REQUIRED_SHEETS - set(workbook.sheetnames))
        if missing:
            raise ValueError(f"workbook is missing required sheet(s): {', '.join(missing)}")
        # Access all required audit/support sheets so the workbook contract is explicit.
        for name in ("Service Summary", "Sources & Audit", "Checks"):
            list(workbook[name].iter_rows(values_only=True))
        services = read_services(workbook["Official Services"])
        branches = read_branches(workbook["Branches 24"], workbook["Capacity Model"])
    finally:
        workbook.close()

    outputs = {
        "services.json": services,
        "service_price_summary.json": summarize_services(services),
        "branch_capacity_scenarios.json": build_capacity_records(branches),
        "branch_revenue_explanations.json": build_revenue_explanations(branches),
        "financial_methodology.json": build_methodology(),
    }
    for filename, payload in outputs.items():
        atomic_write_json(output_dir / filename, payload)
        LOGGER.info("Wrote %s", output_dir / filename)
    return {
        "service_variants": len(services),
        "service_categories": len(outputs["service_price_summary.json"]),
        "branches": len(branches),
        "output_files": len(outputs),
    }


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    try:
        counts = prepare()
    except (FileNotFoundError, ValueError, OSError) as exc:
        LOGGER.error("Financial/services data preparation failed: %s", exc)
        return 1
    LOGGER.info("Validation passed: %s", dict(counts))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
